import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

def update_text_preserving_style(shape, new_text):
    if not shape.has_text_frame:
        return
    tf = shape.text_frame
    if not tf.paragraphs:
        p = tf.add_paragraph()
        p.text = new_text
        return
    
    first_p = tf.paragraphs[0]
    font_name = None
    font_size = None
    font_bold = None
    font_color_rgb = None
    
    if first_p.runs:
        r = first_p.runs[0]
        font_name = r.font.name
        font_size = r.font.size
        font_bold = r.font.bold
        try:
            if r.font.color and r.font.color.type:
                font_color_rgb = r.font.color.rgb
        except Exception:
            pass

    lines = new_text.split('\n')
    tf.clear()
    
    for i, line in enumerate(lines):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        
        run = p.add_run()
        run.text = line
        if font_name:
            run.font.name = font_name
        if font_size:
            run.font.size = font_size
        if font_bold is not None:
            run.font.bold = font_bold
        if font_color_rgb:
            try:
                run.font.color.rgb = font_color_rgb
            except Exception:
                pass

def main():
    src_path = r'C:/Users/BIBEK ROUT/Downloads/Sports Injury Risk Detection From Video.pptx'
    prs = pptx.Presentation(src_path)
    
    print(f"Loaded presentation: {len(prs.slides)} slides.")
    
    # 1. Slide 4 (Workflow)
    s4 = prs.slides[3]
    for shape in s4.shapes:
        if shape.has_text_frame:
            t = shape.text_frame.text.strip()
            if t == "Create account & Sign In":
                update_text_preserving_style(shape, "Role-Based Account & Sign In")
                print("Updated Slide 4: Account creation to Role-Based")
            elif t == "Risk Engine":
                update_text_preserving_style(shape, "XGBoost ML Risk Engine")
                print("Updated Slide 4: Risk Engine to XGBoost ML Risk Engine")
            elif "Video → Pose & Biomechanics → Risk" in t or "Video ? Pose & Biomechanics" in t:
                update_text_preserving_style(shape, "Video → Pose & Biomechanics → XGBoost ML Risk Model → Recommendations → Reports")
                print("Updated Slide 4: Bottom pipeline")

    # 2. Slide 5 (Backend & API)
    s5 = prs.slides[4]
    for shape in s5.shapes:
        if shape.has_text_frame:
            t = shape.text_frame.text.strip()
            if "Registration, login and protected" in t or "Registration, login & Role-Based" in t:
                update_text_preserving_style(shape, "Registration, login & Role-Based Access Control (RBAC)")
                print("Updated Slide 5: Auth description with RBAC")

    # 3. Slide 6 (Frontend & User Interface)
    s6 = prs.slides[5]
    for shape in s6.shapes:
        if shape.has_text_frame:
            t = shape.text_frame.text.strip()
            if "01 . Registration & Login" in t:
                update_text_preserving_style(shape, "01 . Registration & Login\nRole-based secure account creation (Athlete/Coach/Staff).")
                print("Updated Slide 6: Registration & Login feature")

    # 4. Slide 8 (Registration Page Features)
    s8 = prs.slides[7]
    for shape in s8.shapes:
        if shape.has_text_frame:
            t = shape.text_frame.text.strip()
            if "User Registration" in t and "Responsive Design" in t:
                update_text_preserving_style(shape, "Role-Based User Registration\nAthlete & Coach Profiles\nResponsive Design")
                print("Updated Slide 8: User Registration feature")

    # 5. Slide 10 (Login Page Features)
    s10 = prs.slides[9]
    for shape in s10.shapes:
        if shape.has_text_frame:
            t = shape.text_frame.text.strip()
            if "Email and password-based login." in t or "Role-based authentication" in t:
                update_text_preserving_style(shape, "Role-based authentication & login.\nUser-Friendly Input Fields")
                print("Updated Slide 10: Login Features")

    # 6. Slide 21 (Risk Scoring + Recommendations + Testing)
    s21 = prs.slides[20]
    for shape in s21.shapes:
        if shape.has_text_frame:
            t = shape.text_frame.text.strip()
            if t in ["Rule-Based Scoring Engine", "XGBoost ML Risk Engine", "XGBoost Machine Learning Scoring Engine", "Explainable Rule-Based Scoring Engine"]:
                update_text_preserving_style(shape, "XGBoost ML Risk Engine")
                print("Updated Slide 21: Title to XGBoost ML Risk Engine")
            elif "Current implementation is explainable" in t or "Trained XGBoost ML" in t or "scoring engine based on biomechanical" in t:
                update_text_preserving_style(
                    shape, 
                    "Trained XGBoost machine learning model evaluates extracted biomechanical features to predict a 0–100 injury risk score with high explainability."
                )
                print("Updated Slide 21: Description")
            elif t in ["WHY RULE-BASED?", "WHY XGBOOST?", "WHY EXPLAINABLE RULE-BASED?"]:
                update_text_preserving_style(shape, "WHY XGBOOST & SYNTHETIC DATA?")
                print("Updated Slide 21: Header")
            elif "No labeled real-world injury-outcome dataset" in t or "High accuracy on tabular" in t or "Ensures high transparency" in t:
                update_text_preserving_style(
                    shape, 
                    "Addresses clinical data scarcity by training on validated synthetic biomechanical distributions, providing high accuracy, fast inference, and feature importance."
                )
                print("Updated Slide 21: Explanation text")

    # 7. Slide 23 (Results, Limitations & Future Scope)
    s23 = prs.slides[22]
    for shape in s23.shapes:
        if shape.has_text_frame:
            t = shape.text_frame.text.strip()
            if t in ["Secure registration & login", "Role-Based Access Control (RBAC)"]:
                update_text_preserving_style(shape, "Role-Based Access Control (RBAC)")
                print("Updated Slide 23: RBAC result")
            elif t in ["Injury risk score generation", "XGBoost ML injury risk prediction", "Explainable injury risk scoring"]:
                update_text_preserving_style(shape, "XGBoost ML injury risk prediction (0–100)")
                print("Updated Slide 23: Risk score result")
            elif t in ["Rule-based risk engine", "Dataset scaling", "Rule-based baseline scoring"]:
                update_text_preserving_style(shape, "Synthetic training baseline")
                print("Updated Slide 23: Limitation title")
            elif "Current risk scoring is explainable" in t or "Model accuracy scales" in t or "relies on validated biomechanical" in t:
                update_text_preserving_style(shape, "Trained on synthetic biomechanical data; real-world clinical validation is ongoing.")
                print("Updated Slide 23: Limitation text")
            elif "Real labeled injury dataset" in t:
                update_text_preserving_style(shape, "Real-world clinical dataset validation")
                print("Updated Slide 23: Future scope dataset")
            elif "Trained ML risk model" in t or "Deep Learning" in t or "Trained XGBoost" in t:
                update_text_preserving_style(shape, "Deep Learning & Real-time Edge Processing")
                print("Updated Slide 23: Future scope ML")
            elif "Video  →  Analysis  →  Risk" in t or "Video  ?  Analysis" in t or "Video  →  Analysis  →  Explainable" in t:
                update_text_preserving_style(shape, "Video  →  Analysis  →  XGBoost ML Risk (0–100)  →  Recommendations  →  Reports")
                print("Updated Slide 23: Project takeaway flow")

    # 8. Slide 24 (Why You'll Love Using Our Website!)
    s24 = prs.slides[23]
    for shape in s24.shapes:
        if shape.has_text_frame:
            t = shape.text_frame.text.strip()
            if "Easy Athlete Profiling" in t or "Role-Based Athlete" in t:
                update_text_preserving_style(shape, "🏃 Role-Based Athlete & Coach Profiling\n 🎥 Video-Based Movement Analysis\n 🤖 XGBoost ML-Powered Injury Risk Scoring (0–100)\n 💡 Personalized Corrective Recommendations\n 📄 Interactive Reports & History Tracking")
                print("Updated Slide 24: Highlights")

    # 9. Slide 25 (Conclusion)
    s25 = prs.slides[24]
    for shape in s25.shapes:
        if shape.has_text_frame:
            t = shape.text_frame.text.strip()
            if "Our project provides an AI-powered platform" in t:
                new_conclusion = (
                    "Our project provides an AI-powered platform for analyzing athlete movement from video. "
                    "The system features role-based access for athletes and coaches, performs pose estimation and biomechanical analysis, "
                    "and predicts injury risk (0–100) using a trained XGBoost machine learning model. "
                    "It also provides personalized recommendations to help improve movement and reduce potential injury risks. "
                    "In the future, we plan to validate the model against large-scale clinical cohorts and support multi-camera tracking. "
                    "Overall, the project demonstrates a complete, reliable end-to-end AI pipeline for video-based sports injury assessment."
                )
                update_text_preserving_style(shape, new_conclusion)
                print("Updated Slide 25: Conclusion text")

    # Save to both workspace and Downloads with fresh names
    out_workspace = r'd:/Sports_Injury_detection/Sports_Injury_Risk_Detection_XGBoost.pptx'
    out_downloads = r'C:/Users/BIBEK ROUT/Downloads/Sports_Injury_Risk_Detection_XGBoost.pptx'
    
    prs.save(out_workspace)
    prs.save(out_downloads)
    print(f"\nSaved successfully to:\n- {out_workspace}\n- {out_downloads}")

if __name__ == "__main__":
    main()
